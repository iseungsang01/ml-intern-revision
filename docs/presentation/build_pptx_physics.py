# -*- coding: utf-8 -*-
"""물리 근거 덱 — 결과를 받치는 물리를 열 개의 증명과 하나의 시간영역 모형으로 보고한다 (21장, 16:9).

출력: docs/presentation/KSTAR_CES_물리근거.pptx

의사결정 덱(build_pptx_decision.py)의 도형 세트 KIT를 그대로 복사해 쓰며, 장표는 그 세트만 쓴다.
검사: py docs/presentation/check_deck_rules.py docs/presentation/KSTAR_CES_물리근거.pptx

모든 수치의 출처는 docs/theory_results_physics.md (§8bb), docs/theory_foundations.md (§8ba),
docs/theory_plasma_physics.md, THESIS_RESULTS.md §8이다. 모델은 학습·채점하지 않았다.

사용법 (repo root에서):
    py docs/presentation/build_pptx_physics.py
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "KSTAR_CES_물리근거.pptx")

# ============================== 도형 세트 (의사결정 덱과 동일) ==============================
NAVY = RGBColor(0x13, 0x33, 0x5F)
BLUE = RGBColor(0x2B, 0x6C, 0xB0)
TEAL = RGBColor(0x1B, 0x9E, 0x8A)
ORANGE = RGBColor(0xE8, 0x74, 0x3B)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
RED = RGBColor(0xC0, 0x39, 0x2B)
GRAY = RGBColor(0x5B, 0x66, 0x70)
LGRAY = RGBColor(0xE8, 0xEC, 0xF1)
CARDBG = RGBColor(0xF4, 0xF7, 0xFA)
MGRAY = RGBColor(0x9A, 0xA5, 0xB1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x2B, 0x35)

FONT = "Malgun Gothic"
MONO = "Consolas"

KIT = dict(
    W=13.333, H=7.5,
    ML=0.52, CW=12.29,
    TITLE_TOP=0.50, TITLE_H=0.72, TITLE_PT=24.0,
    RULE_Y=1.22, RULE_W=2.20, RULE_T=0.04,
    LEAD_TOP=1.42, LEAD_H=0.50, LEAD_PT=12.5,
    CONTENT_TOP=1.96, SAFE_BOTTOM=6.90,
    FOOT_RULE_Y=6.98, FOOT_Y=7.06, FOOT_PT=9.0,
    GUT=0.20, CARD_H=1.66, CARD_BAR=0.06,
    CARD_TITLE_PT=13.5, CARD_BODY_PT=11.5,
    HEAD_H=0.44, ROW_H=0.47, TABLE_PT=12.0, HEAD_PT=12.0,
    BAND_H=0.78, BAND_PT=13.0,
    ARROW_W=2.60, ARROW_H=0.62, ARROW_GAP=0.16, ARROW_PT=12.0,
    SECTION_NUM_PT=96.0, SECTION_TITLE_PT=24.0,
)
COL3 = (KIT["CW"] - 2 * KIT["GUT"]) / 3.0
COL2 = (KIT["CW"] - KIT["GUT"]) / 2.0
X3 = [KIT["ML"], KIT["ML"] + COL3 + KIT["GUT"], KIT["ML"] + 2 * (COL3 + KIT["GUT"])]
X2 = [KIT["ML"], KIT["ML"] + COL2 + KIT["GUT"]]

prs = Presentation()
prs.slide_width = Inches(KIT["W"])
prs.slide_height = Inches(KIT["H"])
BLANK = prs.slide_layouts[6]
_page = {"n": 0}
FOOTER = "KSTAR CES 나우캐스팅 · 물리 근거 덱 (§8ba · §8bb)"


def _fill(sp, color):
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = color


def rect(s, x, y, w, h, fill=None, line=None, line_w=0.75):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(sp, fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def text(s, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=3, line_spacing=1.12):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(1)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for seg in para:
            t, pt, color, bold, font = (tuple(seg) + (None,) * 5)[:5]
            r = p.add_run()
            r.text = t
            r.font.size = Pt(pt)
            r.font.color.rgb = color if color is not None else DARK
            r.font.bold = bool(bold)
            r.font.name = font or FONT
    return tb


def slide(page=True):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, KIT["W"], KIT["H"], fill=WHITE).name = "BG"
    if page:
        _page["n"] += 1
        rect(s, KIT["ML"], KIT["FOOT_RULE_Y"], KIT["CW"], 0.012, fill=MGRAY).name = "KIT.footrule"
        text(s, KIT["ML"], KIT["FOOT_Y"], KIT["CW"] - 1.0, 0.24,
             [[(FOOTER, KIT["FOOT_PT"], GRAY, False, None)]]).name = "KIT.footer"
        text(s, KIT["ML"] + KIT["CW"] - 1.0, KIT["FOOT_Y"], 1.0, 0.24,
             [[("%d" % _page["n"], KIT["FOOT_PT"], GRAY, False, MONO)]],
             align=PP_ALIGN.RIGHT).name = "KIT.pageno"
    return s


def title_bar(s, title, accent=ORANGE):
    text(s, KIT["ML"], KIT["TITLE_TOP"], KIT["CW"], KIT["TITLE_H"],
         [[(title, KIT["TITLE_PT"], NAVY, True, None)]],
         anchor=MSO_ANCHOR.BOTTOM, line_spacing=1.0).name = "KIT.title"
    rect(s, KIT["ML"], KIT["RULE_Y"], KIT["RULE_W"], KIT["RULE_T"], fill=accent).name = "KIT.rule"


def lead(s, txt):
    text(s, KIT["ML"], KIT["LEAD_TOP"], KIT["CW"], KIT["LEAD_H"],
         [[(txt, KIT["LEAD_PT"], GRAY, False, None)]], space_after=0).name = "KIT.lead"


def card(s, x, y, w, title, lines, accent=NAVY, h=None):
    h = KIT["CARD_H"] if h is None else h
    rect(s, x, y, w, h, fill=CARDBG).name = "KIT.card"
    rect(s, x, y, KIT["CARD_BAR"], h, fill=accent).name = "KIT.cardbar"
    paras = [[(title, KIT["CARD_TITLE_PT"], accent, True, None)]]
    paras += [[(ln, KIT["CARD_BODY_PT"], DARK, False, None)] for ln in lines]
    text(s, x + 0.20, y + 0.13, w - 0.36, h - 0.24, paras, space_after=2).name = "KIT.cardtext"


def band(s, y, title, body, accent=ORANGE, w=None, x=None):
    w = KIT["CW"] if w is None else w
    x = KIT["ML"] if x is None else x
    rect(s, x, y, w, KIT["BAND_H"], fill=LGRAY).name = "KIT.band"
    rect(s, x, y, KIT["CARD_BAR"], KIT["BAND_H"], fill=accent).name = "KIT.bandbar"
    text(s, x + 0.20, y + 0.10, w - 0.36, KIT["BAND_H"] - 0.20,
         [[(title + "  ", KIT["BAND_PT"], accent, True, None),
           (body, KIT["BAND_PT"] - 1.5, DARK, False, None)]]).name = "KIT.bandtext"


def table(s, x, y, col_w, head, rows, size=None, head_size=None, row_h=None,
          emphasis=(), align_right=()):
    size = KIT["TABLE_PT"] if size is None else size
    head_size = KIT["HEAD_PT"] if head_size is None else head_size
    row_h = KIT["ROW_H"] if row_h is None else row_h
    tw = sum(col_w)
    rect(s, x, y, tw, KIT["HEAD_H"], fill=LGRAY).name = "KIT.thead"
    cx = x
    for j, cw in enumerate(col_w):
        a = PP_ALIGN.RIGHT if j in align_right else PP_ALIGN.LEFT
        text(s, cx + 0.10, y + 0.06, cw - 0.20, KIT["HEAD_H"] - 0.12,
             [[(head[j], head_size, NAVY, True, None)]], align=a,
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0).name = "KIT.theadtxt"
        cx += cw
    ry = y + KIT["HEAD_H"]
    for i, row in enumerate(rows):
        if i in emphasis:
            rect(s, x, ry, tw, row_h, fill=CARDBG).name = "KIT.trowhl"
        rect(s, x, ry, tw, 0.010, fill=MGRAY).name = "KIT.trule"
        cx = x
        for j, cell in enumerate(row):
            if isinstance(cell, tuple):
                t, color, bold, font = (cell + (None,) * 4)[:4]
            else:
                t, color, bold, font = cell, None, False, None
            a = PP_ALIGN.RIGHT if j in align_right else PP_ALIGN.LEFT
            text(s, cx + 0.10, ry + 0.05, col_w[j] - 0.20, row_h - 0.10,
                 [[(t, size, color, bold, font)]], align=a,
                 anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0).name = "KIT.tcell"
            cx += col_w[j]
        ry += row_h
    rect(s, x, ry, tw, 0.010, fill=MGRAY).name = "KIT.trule"
    return ry


def arrows(s, y, items, accent=BLUE, x0=None):
    x = KIT["ML"] if x0 is None else x0
    for i, (head, sub) in enumerate(items):
        sp = s.shapes.add_shape(
            MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
            Inches(x), Inches(y), Inches(KIT["ARROW_W"]), Inches(KIT["ARROW_H"]))
        _fill(sp, accent if i % 2 == 0 else NAVY)
        sp.line.fill.background()
        sp.shadow.inherit = False
        sp.name = "KIT.arrow"
        text(s, x + 0.26, y + 0.02, KIT["ARROW_W"] - 0.56, KIT["ARROW_H"] - 0.04,
             [[(head, KIT["ARROW_PT"], WHITE, True, None)]],
             anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0).name = "KIT.arrowtxt"
        text(s, x + 0.26, y + KIT["ARROW_H"] + 0.10, KIT["ARROW_W"] - 0.24, 0.70,
             [[(sub, 10.5, GRAY, False, None)]], line_spacing=1.14).name = "KIT.arrowsub"
        x += KIT["ARROW_W"] + KIT["ARROW_GAP"]


def section(s, num, title, sub, accent=ORANGE):
    rect(s, 0, 0, KIT["W"], KIT["H"], fill=NAVY).name = "SECBG"
    text(s, KIT["ML"] + 0.38, 2.10, 4.0, 1.80,
         [[(num, KIT["SECTION_NUM_PT"], accent, True, MONO)]], line_spacing=1.0).name = "SEC.num"
    text(s, KIT["ML"] + 0.38, 3.90, 11.0, 0.70,
         [[(title, KIT["SECTION_TITLE_PT"], WHITE, True, None)]], line_spacing=1.0).name = "SEC.title"
    rect(s, KIT["ML"] + 0.38, 4.72, KIT["RULE_W"], KIT["RULE_T"], fill=accent).name = "SEC.rule"
    text(s, KIT["ML"] + 0.38, 4.94, 10.5, 0.50,
         [[(sub, 13.0, LGRAY, False, None)]]).name = "SEC.sub"


def note(s, txt):
    s.notes_slide.notes_text_frame.text = txt.strip()
    return s


Y0 = KIT["CONTENT_TOP"]


def after_table(n_rows, row_h=None, gap=0.20):
    row_h = KIT["ROW_H"] if row_h is None else row_h
    return Y0 + KIT["HEAD_H"] + n_rows * row_h + gap


# ================================ 장표 ======================================
def s01():
    s = slide(page=False)
    rect(s, 0, 0, KIT["W"], KIT["H"], fill=NAVY).name = "COVERBG"
    rect(s, KIT["ML"] + 0.38, 2.05, KIT["RULE_W"], KIT["RULE_T"], fill=ORANGE)
    text(s, KIT["ML"] + 0.38, 2.30, 11.5, 0.95,
         [[("결과를 받치는 물리", 40, WHITE, True, None)]], line_spacing=1.0)
    text(s, KIT["ML"] + 0.38, 3.30, 11.5, 0.55,
         [[("KSTAR CES 나우캐스팅 · 열 개의 증명과 하나의 시간영역 모형", 20, ORANGE, True, None)]], line_spacing=1.0)
    text(s, KIT["ML"] + 0.38, 3.95, 11.0, 0.80,
         [[("연구의 모든 수치 결과를 플라즈마의 보존식에서 시간 t의 식으로 유도하고, 641개 방전과 얼린 "
            "TEST 산출물로 각 유도를 검산한 결과를 보고한다. 모델은 학습·채점·선택하지 않았다.",
            14, LGRAY, False, None)]], line_spacing=1.16)
    band(s, 5.20, "이 덱의 결론",
         "두 타깃은 같은 식을 따르되 우변이 다르다. 이온 온도의 우변은 빠른 진단이 보고, "
         "회전의 우변은 데이터에 없다. 그 한 사실이 모든 결과를 정한다.", accent=ORANGE)
    text(s, KIT["ML"] + 0.38, 6.42, 11.5, 0.40,
         [[("이승상 · 서울대학교 원자핵공학과 · 2026-10-13", 12.5, MGRAY, False, None)]])
    return note(s, "결론 한 문장을 먼저 말한다. 뒤의 모든 장은 이 문장을 증명하는 자리다.")


def s02():
    s = slide()
    title_bar(s, "모든 결과는 보존식 우변에 관측 가능한 항이 있는가로 설명된다")
    lead(s, "열 개의 독립 증명이 결과 하나씩을 제1원리에서 유도하고 데이터로 검산하였다. 여섯 카드가 "
            "그 요지이며, 각 카드는 뒤의 한 장으로 받쳐진다.")
    y = Y0
    card(s, X3[0], y, COL3, "① 하나의 식, 두 개의 우변", [
        "dy/dt = -(y-m)/τ + 잡음 + ELM 충격",
        "Tᵢ의 m·충격은 Tₑ·nₑ가 정하고 ECEI·BES가 본다.",
        "V_rot의 것은 토크가 정하며 데이터에 없다.",
    ], accent=NAVY)
    card(s, X3[1], y, COL3, "② Tᵢ 헤드라인은 천이 물리다", [
        "+0.24의 90~114 %가 ELM·천이 행에서 온다.",
        "조용한 절반에서는 보간에 4/4 지고,",
        "상위 5 % 천이 행에서는 4/4 이긴다.",
    ], accent=GREEN)
    card(s, X3[2], y, COL3, "③ V_rot 판정은 토크 부재의 귀결이다", [
        "최적 인과 예측은 이력의 잡음 필터다.",
        "중앙값 하한 0.92 위에 모델 1.14가 앉고",
        "집계 양수는 상위 1 % 행(105행)이 준다.",
    ], accent=ORANGE)
    y2 = y + KIT["CARD_H"] + 0.18
    card(s, X3[0], y2, COL3, "④ 빠른 진단 이득 +0.05는 한계 안이다", [
        "상호정보 선형 한계 +0.047~+0.056.",
        "측정 +0.048이 그 안에 앉는다.",
        "비선형 초과분은 관측되지 않는다.",
    ], accent=BLUE)
    card(s, X3[1], y2, COL3, "⑤ 50 ms는 필터가 아니라 ELM 척도다", [
        "필터 수렴은 2~3 스텝에서 끝나 26~43 %만 설명.",
        "나머지는 ELM 회복 60 ms와 결합 18~39 ms.",
        "자릿수 정합까지만 주장한다.",
    ], accent=TEAL)
    card(s, X3[2], y2, COL3, "⑥ 기록 정정 두 건이 나왔다", [
        "PR2 폴백 43 %는 집계 인공물(실제 0.3~0.5 %).",
        "긴 간극의 V_rot 회복은 빠른 채널이 아니라",
        "양측 라벨 문맥이다. 원장에 주석을 달았다.",
    ], accent=RED)
    band(s, y2 + KIT["CARD_H"] + 0.22, "산출물",
         "문서 3편(foundations · plasma_physics · results_physics) · 증명 스크립트 10폴더 · "
         "원장 §8ba·§8bb · 검산 analyze_foundations.py")
    return note(s, "이 장만 보고 나가도 되게 만든 장이다. 여섯 카드는 각각 뒤의 한 장과 대응한다.")


def s03():
    s = slide()
    title_bar(s, "열 개의 증명은 각각 결과 하나를 유도하고 수치로 검산하였다")
    lead(s, "형식은 같다. 명제 → 유도 → 시간 t의 식 → 데이터 검산 → 기록과의 연결 → 한계와 반증 조건. "
            "검산은 641개 방전 CSV 전수와 얼린 TEST 산출물에서만 하였다.")
    table(s, KIT["ML"], Y0,
          [0.80, 4.10, 4.40, 2.99],
          ["증명", "설명하는 결과", "유도가 내놓은 수치", "검산 상태"],
          [
              [("P01", NAVY, True, MONO), "Tᵢ 헤드라인 +0.24 (4/4)", "천이 항 = 헤드라인의 90~114 %",
               ("항등식 · 4/4", GREEN, True, None)],
              [("P02", NAVY, True, MONO), "V_rot 동률 · 중앙값 -49 %", "하한 0.921 vs 모델 1.14",
               ("정합 · 상관 0", GREEN, True, None)],
              [("P03", NAVY, True, MONO), "문맥 포화 50 ms", "필터 2~3 스텝, τ_rec 60 ms",
               ("자릿수 · 일부 기각", ORANGE, True, None)],
              [("P04", NAVY, True, MONO), "캠페인 드리프트 1.22σ", "shot 간 분산 몫 0.84 / 0.91",
               ("§8n 재현", GREEN, True, None)],
              [("P05", NAVY, True, MONO), "3 keV 컷", "pedestal 상한 1.0~1.8 keV",
               ("런 통계 재현", GREEN, True, None)],
              [("P06", NAVY, True, MONO), "빠른 진단 이득 +0.05", "선형 한계 +0.047~+0.056",
               ("한계 안", GREEN, True, None)],
              [("P07", NAVY, True, MONO), "MNAR · 독립 결측", "결측률 ×5.8, τ_B 26~35 ms",
               ("일부 예측 실패", ORANGE, True, None)],
              [("P08", NAVY, True, MONO), "승률 70 % vs 48 %", "p_obs 1.27배 vs 0.99배",
               ("삼분위 재현", GREEN, True, None)],
          ], row_h=0.45)
    band(s, 6.10, "P09 · P10",
         "MAST 전이 null은 무차원 상사(β 7배·ε 2.75배)와 계기 단위 논증으로, 실시간 등급은 "
         "τ_E·ELM 붕괴·모드 주기의 세 척도로 닫혔다.")
    return note(s, "열 증명 가운데 두 개(P03·P07)는 가설 일부가 기각되었다. 그것을 먼저 말한다.")


def s04():
    s = slide(page=False)
    section(s, "01", "물리 근거", "측정 원리 · 두 보존식 · 하나의 시간영역 모형 · 추정기의 오차 · 50 ms")
    return note(s, "결과를 설명하기 전에 도구를 만든다. 다섯 장이다.")


def s05():
    s = slide()
    title_bar(s, "CES의 두 타깃은 같은 광자 예산에서 나오고 서로 직교한다")
    lead(s, "CES는 C VI 529 nm 선의 Doppler 폭에서 이온 온도를, Doppler 이동에서 회전을 읽는다. "
            "광자 N개를 세는 Poisson 통계가 두 값의 정밀도 하한(Cramér–Rao 한계)을 정한다.")
    table(s, KIT["ML"], Y0,
          [3.40, 3.00, 5.89],
          ["양", "식", "데이터로 확인한 것"],
          [
              ["선폭 → Tᵢ", ("σ_λ = (λ₀/c)·√(kT/m)", DARK, True, MONO),
               "593 eV에서 FWHM 0.287 nm, 열속도 69 km/s"],
              ["회전 정밀도", ("σ_v = v_th / √N", DARK, True, MONO),
               "V_rot 바닥 4.3 km/s → N = 258 광자"],
              ["온도 정밀도", ("σ_T/T = √(2/N)", DARK, True, MONO),
               "Tᵢ 바닥 46.4 eV → N = 326 광자 (비 1.27)"],
              ["중심·폭의 직교성", ("I_μσ = 0", DARK, True, MONO),
               "ML 추정 상관 +0.04 → 두 타깃이 독립으로 결측(일치 72.8 %)"],
              ["느린 이유", ("σ ∝ τ^(-1/2)", DARK, True, MONO),
               "1 km/s 정밀도에는 10 ms 프레임의 16배 적분이 필요하다"],
          ])
    band(s, after_table(5), "읽는 법",
         "서로 다른 타깃에서 따로 잰 두 잡음 바닥이 같은 광자수(약 300/프레임)로 모인다. "
         "타깃의 재현성 한계는 계측 잡음이며, 그 아래로는 어떤 모델도 내려갈 수 없다.", accent=NAVY)
    return note(s, "측정 원리가 문제를 만든다. 광자 통계가 시간 분해능과 정밀도를 교환시킨다.")


def s06():
    s = slide()
    title_bar(s, "두 보존식의 우변 가운데 빠른 진단이 닿는 항은 이온 에너지 쪽뿐이다")
    lead(s, "이온 에너지와 토로이달 각운동량은 값이 아니라 변화율을 구속한다. 각 항을 보유 채널에 대조하면 "
            "라우팅이 설계 선택이 아니라 보존식의 결과임이 드러난다.")
    card(s, X2[0], Y0, COL2, "이온 에너지: (3/2) n dTᵢ/dt = Q_ei(Tₑ-Tᵢ) + S - 손실", [
        "Q_ei ∝ nₑ(Tₑ-Tᵢ): ECEI(Tₑ)·BES(nₑ)가 같은 시각에 본다.",
        "평형화 시간 τ_eq = (mᵢ/2mₑ)τₑ = 8~59 ms (CES 영역).",
        "100 keV 빔은 전자를 71 % 데우므로 Tᵢ는 Tₑ를 따른다.",
    ], accent=GREEN)
    card(s, X2[1], Y0, COL2, "각운동량: d(nmR²ω)/dt = T_NBI + T_NTV + T_int - L/τ_φ", [
        "T_NBI = 2PR_tan/v_b: 파워가 아니라 기하·에너지다.",
        "T_NTV ∝ δB²: 100 Hz Mirnov는 진폭을 71 % 오차로만 담는다.",
        "T_int ∝ ∇Tᵢ: 스칼라만 보유. 우변이 전부 비어 있다.",
    ], accent=ORANGE)
    table(s, KIT["ML"], Y0 + KIT["CARD_H"] + 0.20,
          [2.60, 3.60, 2.60, 3.49],
          ["항", "물리", "데이터에 있는가", "측정 (§8ar · P02)"],
          [
              ["Q_ei(Tₑ-Tᵢ)", "전자–이온 결합", ("ECEI · BES", GREEN, True, None),
               "rate 상관 +0.07, 이력 조건부 partial r 0.26"],
              ["T_NBI", "빔 토크", ("없음", RED, True, None), "T_e~V_rot null (r +0.02)"],
              ["T_NTV", "비축대칭 제동", ("파괴됨", RED, True, None), "파생 MC 특징 전부 음성"],
              ["∇·Π_turb", "난류 운동량 유속", ("BES 요동", ORANGE, True, None),
               "ΔV_rot vs ΔBES 10~200 ms 전부 |r| ≤ 0.014"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    return note(s, "같은 측정이 에너지식에서는 신호를 잡으므로 회전의 null은 방법의 한계가 아니다.")


def s07():
    s = slide()
    title_bar(s, "두 타깃은 하나의 시간영역 식을 따르고 계수의 출처만 다르다")
    lead(s, "보존식을 선형화하면 1차 완화 + 확률 구동 + 사건 충격이 된다. 이 식 하나가 열 증명의 공통 "
            "출발점이며, 표의 계수가 어느 진단에서 오는지가 두 타깃을 가른다.")
    band(s, Y0, "(M1)",
         "dy/dt = -(y-m)/τ + √(2σ_s²/τ)·ξ(t) + Σ_k A_k·ḣ(t-t_k),   h(s) = Θ(s)·exp(-s/τ_rec),   "
         "관측 ỹ(t_j) = y(t_j) + n_j,  Δ = 10 ms", accent=NAVY)
    table(s, KIT["ML"], Y0 + KIT["BAND_H"] + 0.20,
          [2.30, 5.00, 4.99],
          ["계수", "Tᵢ (에너지 보존)", "V_rot (각운동량 보존)"],
          [
              ["완화 τ", "τ_eq ∥ τ_E; 데이터 103~140 ms, IPB98 77 ms", "τ_φ ~ τ_E; 데이터 89 ms → Pr ~ 1"],
              ["평형 m", ("Tₑ·nₑ·P_NBI의 함수 — ECEI·BES가 본다", GREEN, True, None),
               ("토크의 함수 — 데이터에 없다", RED, True, None)],
              ["충격 A_k h", "ELM 붕괴: ms 이하 계단, κ_T ~ 1.03", "토크 램프: κ_V = Δ/2τ_φ = 0.056"],
              ["사건 가시성 p_obs", ("ECEI 급변 동반 농축 1.27~1.34배", GREEN, True, None),
               ("0.99배 = 우연", RED, True, None)],
              ["잡음 σ_n", "39~46 eV, 광자 약 300개", "3~4 km/s, 같은 광자 예산"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    return note(s, "식의 형태는 장치와 무관하고 계수는 전부 장치·계기 고유다. 이것이 P09의 전제이기도 하다.")


def s08():
    s = slide()
    title_bar(s, "추정기의 오차를 시간으로 쓰면 두 항등식이 모든 결과를 담는다")
    lead(s, "사건 시각 t_k를 기준으로 이력 필터·양측 보간·나우캐스터의 오차를 쓰고 짝지은 차이를 시간 적분하면 "
            "방전별 승패 식이, 제곱오차 비중으로 쓰면 skill의 정확한 분해가 나온다.")
    band(s, Y0, "이력 필터 (간극 g)",
         "E[e²](g) = e^(-2g/τ)·P + σ_s²(1-e^(-2g/τ)) + σ_n²  — 양측 교량은 조용한 구간에서 1.5σ_n², "
         "사건 행에서 A_k/2를 틀린다.", accent=NAVY)
    band(s, Y0 + KIT["BAND_H"] + 0.18, "(M6) 방전별 짝 차이",
         "D = p_obs · r · N_e κ A²/n - c_b   — 사건이 보이고(p_obs) 크기가 회수되며(r) 충격 에너지가 "
         "본류 비용 c_b를 넘을 때만 이긴다.", accent=TEAL)
    band(s, Y0 + 2 * (KIT["BAND_H"] + 0.18), "(M7) skill의 정확한 분해",
         "skill = f_I · r_I - (1-f_I) · ℓ_B   — 천이 행의 기준선 SSE 비중 × 그 행의 회수율, 빼기 본류 손실. "
         "4분할에서 1e-16 정확도로 성립.", accent=ORANGE)
    y = Y0 + 3 * (KIT["BAND_H"] + 0.18)
    card(s, X2[0], y, COL2, "이 두 식이 설명하는 것", [
        "헤드라인(P01) · 승률 비대칭(P08) · V_rot 꼬리 부등식(P02)",
        "· 지연 비용(P10)이 전부 이 두 식의 특수한 경우다.",
    ], accent=NAVY, h=1.20)
    card(s, X2[1], y, COL2, "본류 비용 c_b는 구동원과 무관하다", [
        "정보 없는 입력 d개는 추정 분산만 더한다(σ²d/n).",
        "조용 방전 삼분위 손실 -0.072(Tᵢ) = -0.073(V_rot).",
    ], accent=GRAY, h=1.20)
    return note(s, "P08의 D와 P01의 항등식이 이 덱의 두 핵심 식이다. 뒤의 결과 장은 이 둘의 적용이다.")


def s09():
    s = slide()
    title_bar(s, "50 ms 포화에서 필터 수렴 가설은 기각되고 ELM 회복 척도가 남는다")
    lead(s, "세 후보를 따로 재었다. 잡음 Markov 필터의 냉시동 수렴, 전자–이온 평형화 시간, ELM 뒤 pedestal "
            "회복 시간이다. 측정된 포화는 5 스텝(50 ms)이다.")
    table(s, KIT["ML"], Y0,
          [3.30, 3.30, 2.90, 2.79],
          ["후보 척도", "유도가 준 값", "측정 50 ms와의 관계", "판정"],
          [
              ["필터 냉시동 수렴 (ρ(1-K))^2k", "2 % 포화 2~3 스텝", "20 ms 결손의 26~43 %만",
               ("주원인 기각", RED, True, None)],
              ["전자–이온 평형화 τ_eq", "18~39 ms (CES 구석)", "같은 자릿수", ("정합", GREEN, True, None)],
              ["ELM 회복 τ_rec (2,783건)", "60 ms (정의별 55~64)", "같은 자릿수", ("정합", GREEN, True, None)],
              ["회복 형태 식별 (CRLB)", "문맥 ~ 1.5 τ_rec 필요", "정보 논증", ("정합", TEAL, True, None)],
          ], row_h=0.45)
    band(s, after_table(4, 0.45), "판정",
         "50 ms는 τ_eq와 τ_rec 사이에 놓이며 같은 pedestal 물리라는 진술은 자릿수 수준으로만 방어한다. "
         "P10의 물리 크기 충격 τ_rec 25~37 ms, P07의 배경 완화 τ_B 26~35 ms도 같은 자릿수다.", accent=TEAL)
    return note(s, "가설 하나가 기각된 장이다. 기각을 숨기지 않는 것이 이 덱의 신뢰 근거다.")


def s10():
    s = slide(page=False)
    section(s, "02", "결과의 설명", "Tᵢ 헤드라인 · V_rot 판정 · 승률 비대칭 · 정보 한계 · 계측 · 결측 · 전이 · 실시간",
            accent=TEAL)
    return note(s, "도구로 결과를 설명한다. 일곱 장이다.")


def s11():
    s = slide()
    title_bar(s, "Tᵢ 헤드라인은 천이 행이 만들며 조용한 행에서는 보간에 진다", accent=TEAL)
    lead(s, "같은 pedestal 붕괴가 Tᵢ·Tₑ·nₑ를 동시에 떨어뜨린다. 양측 보간은 점프를 가로질러 A/2를 틀리고 "
            "동시 관측을 보는 인과 모델은 (1-κ)A로 줄인다. 이것이 보간을 이길 수 있는 유일한 자리다.")
    table(s, KIT["ML"], Y0,
          [1.40, 2.20, 2.20, 2.20, 2.20, 2.09],
          ["분할", "헤드라인", "f_I (천이 SSE 비중)", "r_I (회수율)", "천이 항 f_I·r_I", "비중"],
          [
              [("42", DARK, True, MONO), ("+0.174", DARK, False, MONO), ("0.734", DARK, False, MONO),
               ("+0.271", DARK, False, MONO), ("+0.199", GREEN, True, MONO), ("1.14", DARK, False, MONO)],
              [("1", DARK, True, MONO), ("+0.248", DARK, False, MONO), ("0.723", DARK, False, MONO),
               ("+0.309", DARK, False, MONO), ("+0.223", GREEN, True, MONO), ("0.90", DARK, False, MONO)],
              [("7", DARK, True, MONO), ("+0.257", DARK, False, MONO), ("0.734", DARK, False, MONO),
               ("+0.333", DARK, False, MONO), ("+0.244", GREEN, True, MONO), ("0.95", DARK, False, MONO)],
              [("123", DARK, True, MONO), ("+0.264", DARK, False, MONO), ("0.779", DARK, False, MONO),
               ("+0.340", DARK, False, MONO), ("+0.265", GREEN, True, MONO), ("1.00", DARK, False, MONO)],
          ])
    y = after_table(4)
    card(s, X2[0], y, COL2, "조용한 절반 대 천이 상위 5 % (비인과 GP 대비)", [
        "조용한 절반: -0.31 / -0.09 / -0.15 / -0.16 (4/4 패)",
        "상위 5 %: +0.05 / +0.09 / +0.09 / +0.05 (4/4 승)",
        "이력만의 인과 GP는 두 층 모두 진다.",
    ], accent=GREEN)
    card(s, X2[1], y, COL2, "동시성의 직접 측정 (641파일, 216,895쌍)", [
        "r(ΔTᵢ, ΔECEI) 전체 +0.142 → 큰 점프 조건부 +0.238",
        "신뢰구간이 분리되고 lag 0에서만 산다.",
        "shot별 ρ(skill, 천이 비중) +0.27~+0.54, 4/4 유의.",
    ], accent=NAVY)
    return note(s, "헤드라인을 '꼬리에서 온다'(§8ax)에서 '천이 물리에서 온다'로 한 단계 더 물리로 내린 장이다.")


def s12():
    s = slide()
    title_bar(s, "V_rot의 모든 판정은 토크가 비관측일 때의 이력 필터 식에서 나온다", accent=TEAL)
    lead(s, "운동량 균형의 해 V(t) = ∫h(t-s)T(s)ds에서 토크가 비관측이면 최적 인과 예측은 이력만의 "
            "잡음 OU 필터다. 그 중앙값 하한과 양측 교량의 하한을 측정과 나란히 둔다.")
    table(s, KIT["ML"], Y0,
          [4.00, 2.60, 2.60, 3.09],
          ["1-step 중앙값 오차 (persistence = 1)", "유도된 하한", "측정 (s42 컷)", "읽기"],
          [
              ["최적 인과 필터", ("0.921", NAVY, True, MONO), "—", "어떤 인과 방법도 아래로 못 간다"],
              ["나우캐스터(MSE 학습)", "≥ 0.921", ("1.140", ORANGE, True, MONO), "하한 위 24 % (헤지 비용)"],
              ["인과 GP", "≥ 0.921", ("1.009", DARK, False, MONO), "하한 위"],
              ["양측 교량 (PCHIP)", ("0.798", NAVY, True, MONO), ("0.791", GREEN, True, MONO),
               "1 % 안에서 일치"],
          ], row_h=0.45)
    y = after_table(4, 0.45)
    card(s, X2[0], y, COL2, "집계 양수는 꼬리 부등식 G_T > D_B가 준다", [
        "상위 1 %(105행)가 PCHIP SSE의 83 %, skill 전부.",
        "G_T +0.414 - D_B +0.024 = +0.390 (항등식).",
        "그래서 PR4는 꼬리가 큰 분할(s42)만 유의하다.",
    ], accent=ORANGE)
    card(s, X2[1], y, COL2, "구동이 비관측임을 모든 창에서 쟀다", [
        "ΔV_rot vs ΔBES·ΔECEI·ΔMC: lag ±2, 창 10~200 ms",
        "전부 |r| ≤ 0.014 (대조 Tᵢ는 +0.15~+0.28).",
        "사건 시각조차 빠른 채널에 비치지 않는다.",
    ], accent=RED)
    return note(s, "V_rot의 네 판정이 전부 한 식의 귀결이라는 것이 이 장이다. 실패가 아니라 식의 결과다.")


def s13():
    s = slide()
    title_bar(s, "방전별 승률 비대칭은 사건이 빠른 채널에 보이는 확률로 설명된다", accent=TEAL)
    lead(s, "D = p_obs·r·E_imp - c_b에서 승률은 방전의 타깃 산포에 증가하되 천장이 p_obs·r > 0인 방전의 "
            "비율로 정해진다. Tᵢ의 천장은 1 근처, V_rot는 0.5~0.6이다.")
    table(s, KIT["ML"], Y0,
          [3.60, 2.40, 2.40, 3.89],
          ["양 (4분할 풀링, 인과 GP 대비)", "Tᵢ", "V_rot", "읽기"],
          [
              ["승률 조용 / 중간 / 변동 삼분위", ("43 / 85 / 86 %", GREEN, True, MONO),
               ("36 / 45 / 59 %", ORANGE, True, MONO), "§8an의 42/83/85 · 34/48/55 재현"],
              ["조용 삼분위 중앙 skill", ("-0.072", DARK, False, MONO), ("-0.073", DARK, False, MONO),
               "본류 비용 -c_b가 구동원과 무관"],
              ["충격행 승률 / 본류행 승률", ("0.744 / 0.594", GREEN, True, MONO),
               ("0.425 / 0.530", RED, True, MONO), "V_rot는 충격행에서 진다 (r ≤ 0)"],
              ["p_obs 대리: 충격행 ECEI 급변 농축", ("1.27배", GREEN, True, None), ("0.99배", RED, True, None),
               "사건 가시성의 직접 측정"],
              ["logit(승) ~ 산포 + p_obs 계수", ("+1.92 ± 0.53", GREEN, True, MONO),
               ("-0.55 ± 0.44", GRAY, False, MONO), "산포를 통제해도 Tᵢ만 가른다"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    band(s, after_table(5, 0.45), "따름정리",
         "Tᵢ의 ELM 붕괴는 계단이라 κ_T ~ 1.03인데 V_rot의 토크 변화는 램프라 κ_V = Δ/2τ_φ = 0.056이다. "
         "같은 진폭이라도 회수 가능한 에너지가 18분의 1이다.", accent=TEAL)
    return note(s, "§8an이 '절반의 이유'로 남긴 잔여를 p_obs·r·κ의 곱으로 메운 장이다.")


def s14():
    s = slide()
    title_bar(s, "빠른 진단이 더하는 정보는 상호정보로 유계이고 측정값은 그 안에 있다", accent=TEAL)
    lead(s, "최선의 이력 예측기 위에 빠른 채널 z를 더해 얻는 최대 이득은 Δskill = r²(y,z|h)·MSE_h/MSE_pers다. "
            "올바른 조건부 변수는 level도 rate도 아니라 필터의 혁신이다.")
    table(s, KIT["ML"], Y0,
          [5.00, 2.70, 4.59],
          ["양", "값", "읽기"],
          [
              ["측정: 모델 - 인과 GP (1-step, s42)", ("+0.048", GREEN, True, MONO), "함의 partial r 0.26, I = 0.036 nat"],
              ["선형 한계, 50 ms 이력 조건부, 빠른 채널 (t, t-1)", ("+0.047", NAVY, True, MONO),
               "측정이 바로 그 위에 앉는다"],
              ["선형 한계, 50 ms 빠른 맥락까지", ("+0.056", NAVY, True, MONO), "비선형 초과분 없음"],
              ["§8ar의 rate 상관 0.07이 주는 한계", ("+0.003", RED, True, MONO), "측정의 1/12 — 올바른 입력이 아니다"],
              ["구조적 천장 K·MSE_h/MSE_pers", ("0.23~0.40", GRAY, False, MONO), "어떤 정보원도 넘지 못한다"],
              ["Mirnov(100 Hz)의 기여", ("0.000", RED, True, MONO), "어느 조건부에서도 0"],
          ], row_h=0.45)
    band(s, after_table(6, 0.45), "경로",
         "이득은 rate가 아니라 잡음 앵커의 탈잡음(level 정보로 nₜ₋₁을 식별)이 운반하고, 그 81~83 %는 "
         "상위 10 % 혁신 행에 산다. 백본은 선형 판독기가 뽑을 만큼을 뽑고 그 이상은 뽑지 못한다.", accent=BLUE)
    return note(s, "용량 축이 평평했던 이유의 정보론 판본이다. 뽑을 선형 정보가 다 뽑혔다.")


def s15():
    s = slide()
    title_bar(s, "캠페인 드리프트는 계측 이득이고 3 keV 초과는 압력 한계 밖이다", accent=TEAL)
    lead(s, "빠른 진단은 x = g(D)·s(t) + c(D), 즉 날짜 D의 계측 이득 곱 플라즈마 상태다. 타깃은 파장비라 D에 "
            "무관하다. 3 keV는 pedestal이 peeling–ballooning 한계 안에서 낼 수 없는 온도다.")
    card(s, X2[0], Y0, COL2, "P04 캠페인 드리프트 (641파일)", [
        "shot 간 분산 몫: BES 0.84 · ECEI 0.91 vs Tᵢ 0.49 · V_rot 0.28",
        "§8n의 1.218 / 0.529 / 0.115 σ를 소수 셋째 자리까지 재현.",
        "shot 내 z-score가 g, c를 정확히 소거; 절대 수준의",
        "추가 정보는 R² 1.4~2.6 %뿐(§8s '무비용'의 정체).",
    ], accent=BLUE, h=2.05)
    card(s, X2[1], Y0, COL2, "P05 3 keV 컷 (PB 압력 한계)", [
        "p_ped,max 4.3~9.6 kPa → Tᵢ,ped,max 1.0~1.8 keV.",
        "3 keV에는 β_pol,ped 1.5~4.5가 필요(상한 0.89).",
        "스파이크 중앙값 5.3 keV의 압력 = PB의 3.06배.",
        "런 통계 1,197행 / 951런 / 85.5 % 단일행 재현.",
    ], accent=ORANGE, h=2.05)
    table(s, KIT["ML"], Y0 + 2.05 + 0.20,
          [4.20, 2.60, 5.49],
          ["검정", "값", "뜻"],
          [
              ["온라인 EWMA 표준화의 편향 (τ_w = 1 s)", ("0.29~0.33 σ", DARK, False, MONO),
               "유도값 b·τ_w = 0.31 / 0.27 σ와 일치 — 배치형의 가격"],
              ["스파이크 행의 V_rot 점프 (Tᵢ 대비)", ("0.05배 vs 5.8배", DARK, False, None),
               "폭 전용 적합 실패의 지문 (p = 0.75)"],
              ["컷 문턱 2.5 → 4 keV에서 이동하는 행", ("0.15~0.19 %", DARK, False, MONO),
               "물리 상한이 2.5 keV 아래라 문턱이 둔감"],
          ], row_h=0.45)
    return note(s, "두 결과 모두 플라즈마가 아니라 계측과 압력 한계가 설명한다는 장이다.")


def s16():
    s = slide()
    title_bar(s, "결측은 광자 예산의 상태 의존이고 실시간 등급은 플라즈마 척도가 정한다", accent=TEAL)
    lead(s, "채택 지시자는 P(obs | N(t), B(t))이며 N은 빔 감쇠·탄소 밀도, B는 ELM 뒤 배경이다. 지연 ℓ의 "
            "가격은 한 격자 스텝 손실의 ℓ/Δ이고, 세 제어 등급은 τ_E·ELM 붕괴·모드 주기로 갈린다.")
    table(s, KIT["ML"], Y0,
          [4.60, 3.00, 4.69],
          ["P07 결측 (641파일, 247,207행)", "값", "읽기"],
          [
              ["타깃 점프 5분위 결측률 (Tᵢ)", ("0.77 → 4.45 %", ORANGE, True, MONO), "×5.8, BES 층화 후에도 ×3.6~7.8"],
              ["BES 수준 최하 / 최상 5분위 결측률", ("19.9 / 3.7 %", ORANGE, True, MONO), "광자 예산의 대리가 결측을 예측"],
              ["BES 급락 뒤 결측 초과분 감쇠 τ_B", ("26~35 ms", DARK, False, MONO), "B(t) = B₀ + B₁e^(-(t-t₀)/τ_B)"],
              ["마스크 일치 / φ", ("72.77 % / 0.041", DARK, False, MONO), "§8at 재현, 독립보다 약간 위"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    y = after_table(4, 0.45)
    card(s, X2[0], y, COL2, "P10 세 지연 등급", [
        "A 수송 제어(τ_E 80~140 ms): 나우캐스터 p99 1.5~9 ms 전부 해당.",
        "B ELM 트리거(<1 ms)·C 모드 제어(0.1~1 ms): 라벨의",
        "10 ms 적분이 배제하므로 1 ms 판정 보류와 무관.",
    ], accent=BLUE, h=1.40)
    card(s, X2[1], y, COL2, "지연의 가격", [
        "skill -0.023/ms(강건)·-0.009/ms(고전), eV로는 1.5 ms에 +1.5 eV.",
        "seq_v2 한 스텝 353,920 MAC · 111 디스패치 → 바닥 0.22~0.33 ms.",
    ], accent=GRAY, h=1.40)
    return note(s, "MNAR을 라벨이 아니라 N(t)·B(t)의 식으로 쓴 장이다. ECEI 증분이 결측을 예측하지 못한 것도 기록했다.")


def s17():
    s = slide()
    title_bar(s, "MAST 사전학습이 아무것도 옮기지 못한 것은 무차원 상사의 결과다", accent=TEAL)
    lead(s, "수송은 (ρ*, ν*, β, q, ε, κ)의 함수다. 두 장치는 다른 점에 있고, 모델 입력은 계기 단위라 "
            "옮겨질 수 있는 것은 식 (M1)의 형태뿐이다. 그 형태는 한 shot이면 KSTAR 데이터가 정한다.")
    table(s, KIT["ML"], Y0,
          [2.60, 2.20, 2.60, 4.89],
          ["무차원 양", "KSTAR", "MAST", "뜻"],
          [
              ["ε = a/R", ("0.278", DARK, False, MONO), ("0.765", DARK, False, MONO), "포획 분율 0.66 vs 0.93"],
              ["β", ("0.45 %", DARK, False, MONO), ("3.2 %", ORANGE, True, MONO), "7배 — 전자기 난류 영역"],
              ["ρ*", ("0.0056", DARK, False, MONO), ("0.0099", DARK, False, MONO), "1.6~2.0배"],
              ["IPB98 τ_E", ("77 ms", DARK, False, MONO), ("48 ms", DARK, False, MONO), "10 ms AR(1) 계수 0.878 vs 0.813"],
              ["전이 가중치 비중", ("1,823,490의 98.2 %", DARK, False, None), "—", "설명할 구조는 식 한 줄(계수 7개)"],
              ["한 중앙 shot (3.98 s)의 SE(ρ)", ("0.027", GREEN, True, MONO), "—", "MAST prior 편향 0.064 = 2.4 SE"],
          ], row_h=0.45)
    band(s, after_table(6, 0.45), "§8az 재현",
         "같은 shot 라벨 3→10에서 네 모델군이 똑같이 +0.26 오르고(퍼짐 0.006), 사전학습 대 무작위 초기화 차는 "
         "0.004~0.013이다. 데이터가 제한하는 양은 prior가 아니라 계기 사상이다.", accent=NAVY)
    return note(s, "장치 간 전이를 다시 시도하지 않을 근거를 무차원수로 적은 장이다.")


def s18():
    s = slide(page=False)
    section(s, "03", "정정과 측정", "기록 정정 두 건 · 유도가 지목한 측정 · 산출물 지도", accent=ORANGE)
    return note(s, "증명이 기록에서 고친 것과 다음에 재야 할 것이다.")


def s19():
    s = slide()
    title_bar(s, "증명 과정에서 원장의 두 문장을 정정하였고 하나를 고쳐 읽었다")
    lead(s, "정정은 모두 얼린 산출물과 코드에서 직접 확인한 뒤 원장 해당 절에 주석으로 달았다. 판정(PR4)과 "
            "주 모델은 바뀌지 않는다.")
    table(s, KIT["ML"], Y0,
          [2.00, 4.20, 6.09],
          ["자리", "원래 문장", "확인된 사실"],
          [
              [("§8ab 블록 2", NAVY, True, MONO), "V_rot 채점 행의 40~44 %가 PR2 폴백",
               ("eval_seq.py의 n_total이 비채점 행까지 셈. 채점 행 기준 0.3~0.5 %", RED, True, None)],
              [("§8az", NAVY, True, MONO), "긴 간극 안에서 빠른 채널이 V_rot를 회복",
               ("대조군 bilstm의 V_rot 분기는 빠른 채널을 보지 않음. 양측 라벨 문맥 효과", RED, True, None)],
              [("§8ar", NAVY, True, MONO), "Tᵢ 결합 크기 = rate 상관 +0.07",
               ("한계의 올바른 입력은 이력 조건부 partial r 0.26 (고쳐 읽기)", ORANGE, True, None)],
          ], row_h=0.50)
    y = after_table(3, 0.50)
    card(s, X2[0], y, COL2, "왜 생겼는가", [
        "첫째는 분모 집계 범위의 차이, 둘째는 대조군의 라우팅을",
        "본문에서 다시 확인하지 않은 것, 셋째는 조건부 변수의 선택이다.",
        "셋 다 유도가 '어떤 양을 재야 하는가'를 먼저 정했기에 드러났다.",
    ], accent=NAVY)
    card(s, X2[1], y, COL2, "무엇이 안 바뀌는가", [
        "Tᵢ 4/4+4/4, 캠페인 4/4, 인과 GP 대비 4/4, V_rot '동률' 판정,",
        "백본 seq_v2, 두 모집단 규칙. 정정은 해석 문장에 한정된다.",
    ], accent=GREEN)
    return note(s, "정정을 숨기지 않고 먼저 말한다. 판정은 그대로다.")


def s20():
    s = slide()
    title_bar(s, "유도가 지목하는 측정은 식의 어느 항을 닫는지까지 정해져 있다")
    lead(s, "음성 결과는 뒤집을 측정을 함께 적어야 결론이 된다. 열 증명이 지목한 측정을 식의 항과 반증 조건으로 "
            "정리하였다. 전부 재학습이 아니라 입력·채점의 문제다.")
    table(s, KIT["ML"], Y0,
          [3.70, 2.90, 1.50, 4.19],
          ["측정", "닫는 항", "지목", "뒤집힐 결과"],
          [
              ["D_α / ELM 지시자", "사건 시각 t_k", ("P01·07·08", GRAY, False, MONO), "p_obs가 1에 가까워짐"],
              ["kHz Mirnov 모드 주파수·모드수", "V_rot Doppler v = 2πRf/n", ("P02·08·10", GRAY, False, MONO),
               "V_rot 충격행 승률 역전"],
              ["Mirnov 대역 파워", "T_NTV ∝ δB²", ("P02", GRAY, False, MONO), "간극 안 토크 일부 관측"],
              ["NBI · RMP 파형", "T_NBI · T_NTV", ("P02·08·09", GRAY, False, MONO), "운동량식 우변이 채워짐"],
              ["CES 프레임별 적합 품질", "채택 지시자 m_j", ("P05·07", GRAY, False, MONO), "값 컷 → 품질 컷, 모집단 병합"],
              ["CES 자속면의 보정 Tₑ", "Q_ei 항", ("P06·03", GRAY, False, MONO), "결합항 부호 0.46 → 일관 양수"],
              ["OU/Kalman 인과 팔 채점", "최적 이력 기준", ("P02·06", GRAY, False, MONO), "빠른 진단 귀속 +0.027~0.048 확정"],
              ["같은 추정기 EWMA 표준화 팔", "g(D), c(D) 온라인 소거", ("P04", GRAY, False, MONO), "첫 1~2 s 워밍업만 남는지"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    return note(s, "여덟 측정 모두 재학습 없이 입력과 채점으로 시작할 수 있다. 데이터 획득 쪽이 다섯이다.")


def s21():
    s = slide()
    title_bar(s, "산출물은 문서 세 편과 증명 폴더 열 개이며 전부 저장소에 있다", accent=GRAY)
    lead(s, "읽는 순서는 결과 문서의 0장(통일 모형과 결과 지도) → 1장(증명 전문) → 2장(종합)이다. "
            "모든 수치는 스크립트 출력에서 전사되었고 커밋으로 고정되어 있다.")
    table(s, KIT["ML"], Y0,
          [4.60, 7.69],
          ["파일 · 절", "내용"],
          [
              [("docs/theory_results_physics.md", NAVY, True, MONO),
               "하나의 모형(M1~M7) · 열 증명 전문 · 종합(관련성 서사 · 정정 · 지목 측정)"],
              [("docs/theory_foundations.md", NAVY, True, MONO),
               "측정 물리 · 추정 이론 · 평가 통계의 유도 17절과 검산"],
              [("docs/theory_plasma_physics.md", NAVY, True, MONO),
               "플라즈마 거동 물리 12장 · CES를 예측하는 두 경로 · Pr ~ 1 · EHO Doppler"],
              [("ces_prediction/experiments/proofs/p01~p10", NAVY, True, MONO),
               "증명 스크립트와 JSON 출력, README"],
              [("ces_prediction/analyze_foundations.py", NAVY, True, MONO),
               "유도 검산 Part A~F → data/.foundations.json"],
              [("THESIS_RESULTS.md §8ba · §8bb", NAVY, True, MONO),
               "원장 기록, §8ab·§8az 정정 주석; PROJECT_KNOWLEDGE.md 교훈"],
          ], size=11.5, head_size=11.5, row_h=0.45)
    band(s, after_table(6, 0.45), "커밋",
         "7012f48 · 052f937 (근거 문서) → 51f72ae · 78d72dc · ad4c803 · 0d66804 (열 증명과 정정). "
         "테스트 52개 통과. 논문 tex·pdf와 기존 발표 덱은 건드리지 않았다.", accent=GRAY)
    return note(s, "어디에 무엇이 있는지로 덱을 닫는다.")


def main():
    for fn in (s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14,
               s15, s16, s17, s18, s19, s20, s21):
        fn()
    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    main()
